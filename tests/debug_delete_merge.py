"""delete_rows가 병합 셀 값을 어떻게 처리하는지 확인."""
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
# R1: 병합 C1:D1 값 '(1) 설계'
ws.cell(1, 3).value = "(1) 설계"
ws.merge_cells("C1:D1")
# R2: 일반
ws.cell(2, 3).value = "0)"
ws.cell(2, 4).value = "Engineering"
# R3: 일반
ws.cell(3, 3).value = "1)"
ws.cell(3, 4).value = "기구설계자(특급)"
# R4: 일반
ws.cell(4, 3).value = "2)"
ws.cell(4, 4).value = "기구설계자(고급)"

print("=== BEFORE ===")
for r in range(1, 5):
    print(f"R{r}: C={ws.cell(r,3).value!r} D={ws.cell(r,4).value!r}")
print("merges:", [str(m) for m in ws.merged_cells.ranges])

# R2 삭제 (병합 해제 없이)
print("\n=== delete_rows(2) WITHOUT unmerge ===")
ws.delete_rows(2, 1)
for r in range(1, 5):
    print(f"R{r}: C={ws.cell(r,3).value!r} D={ws.cell(r,4).value!r}")
print("merges:", [str(m) for m in ws.merged_cells.ranges])

# 다시 만들어서 병합 해제 후 삭제
wb2 = Workbook()
ws2 = wb2.active
ws2.cell(1, 3).value = "(1) 설계"
ws2.merge_cells("C1:D1")
ws2.cell(2, 3).value = "0)"
ws2.cell(2, 4).value = "Engineering"
ws2.cell(3, 3).value = "1)"
ws2.cell(3, 4).value = "기구설계자(특급)"
ws2.cell(4, 3).value = "2)"
ws2.cell(4, 4).value = "기구설계자(고급)"

print("\n=== delete_rows(2) WITH unmerge first ===")
for m in list(ws2.merged_cells.ranges):
    ws2.unmerge_cells(str(m))
ws2.delete_rows(2, 1)
for r in range(1, 5):
    print(f"R{r}: C={ws2.cell(r,3).value!r} D={ws2.cell(r,4).value!r}")
print("merges:", [str(m) for m in ws2.merged_cells.ranges])
