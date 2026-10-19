"""비율조정 로직.xlsx 구조 분석."""
from openpyxl import load_workbook

PATH = r"D:\0_AX 가속화 팀\팀과제\재료비 자동산출\비율조정 로직.xlsx"
wb = load_workbook(PATH)
print("시트:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n=== {name} (max_row={ws.max_row}, max_col={ws.max_column}) ===")
    for r in range(1, ws.max_row + 1):
        vals = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is not None:
                vals.append(f"{ws.cell(r,c).coordinate}={v!r}")
        if vals:
            print(f"R{r}: " + " | ".join(vals))
    print("병합:", sorted(str(m) for m in ws.merged_cells.ranges))
