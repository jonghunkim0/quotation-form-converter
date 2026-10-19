"""출력 결과 검증"""
import openpyxl
import pandas as pd

# ── 표준 Excel 확인 ──
print("=" * 60)
print("=== 표준_견적서_풍산_AZS.xlsx ===")
wb = openpyxl.load_workbook(r"c:\AI_Agent\quotation-form-converter\data\output\표준_견적서_풍산_AZS.xlsx")
print("시트:", wb.sheetnames)
ws = wb["2.상세 견적"]
print(f"max_row={ws.max_row}, max_col={ws.max_column}")

print("\n--- 헤더 + 첫 15행 ---")
for r in range(1, 16):
    vals = []
    for c in range(1, 12):
        v = ws.cell(row=r, column=c).value
        if v is not None:
            col = openpyxl.utils.get_column_letter(c)
            vs = str(v)
            if vs.startswith("="):
                vs = f"[{vs[:20]}]"
            vals.append(f"{col}={vs[:20]}")
    if vals:
        print(f"R{r}: " + " | ".join(vals))

# 소계/합계 행 확인
print("\n--- 소계/합계 행 ---")
for r in range(1, ws.max_row + 1):
    c3 = ws.cell(row=r, column=3).value
    a1 = ws.cell(row=r, column=1).value
    label = c3 or a1
    if label and any(k in str(label) for k in ["소계", "합계", "Sub", "TOTAL", "일반관리", "기업이윤"]):
        j = ws.cell(row=r, column=10).value
        print(f"R{r}: {str(label)[:25]} | J={str(j)[:30] if j else ''}")

wb.close()

# ── 상세 검증 파일 확인 ──
print("\n" + "=" * 60)
print("=== 상세_검증파일.csv ===")
df = pd.read_csv(r"c:\AI_Agent\quotation-form-converter\data\output\상세_검증파일.csv", encoding="utf-8-sig")
print("컬럼:", list(df.columns))
print("행 수:", len(df))
print("\n검증결과 분포:")
print(df["검증결과"].value_counts().head(10).to_string())
print("\n미분류 항목:")
print(df[df["매핑섹션"] == "미분류"].to_string())
