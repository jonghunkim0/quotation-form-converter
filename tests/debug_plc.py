"""PLC 항목 배치 및 _구분 값 디버그"""
import json
import openpyxl
import glob
import os

# 1) 디버그 덤프 확인
print("=== _debug_export.json ===")
try:
    d = json.load(open("data/_debug_export.json", "r", encoding="utf-8"))
    print(f"items: {len(d['items'])}")
    print(f"section_counts: {json.dumps(d['section_counts'], ensure_ascii=False)}")
    # PLC 관련 항목
    for it in d["items"]:
        if "PLC" in it["name"].upper() or "INVERTER" in it["name"].upper():
            print(f"  row={it['row']} name={it['name']!r} section={it['section']!r} bunryu={it['bunryu']!r}")
except Exception as e:
    print(f"  (없음: {e})")

# 2) 최신 출력 파일에서 PLC 항목 위치 확인
print("\n=== 최신 출력 파일 ===")
files = sorted(glob.glob("data/output/*.xlsx"), key=os.path.getmtime)
if files:
    f = files[-1]
    print(f"file: {f} ({os.path.getmtime(f):.0f})")
    wb = openpyxl.load_workbook(f)
    ws = wb["2.상세 견적"]
    # PLC/Inverter/Touch Panel 항목 찾기
    for r in range(5, ws.max_row + 1):
        d_val = str(ws.cell(r, 4).value or "")
        if any(kw in d_val.upper() for kw in ("PLC", "INVERTER", "TOUCH PANEL", "SAFETY")):
            a = ws.cell(r, 1).value
            b = ws.cell(r, 2).value
            c = ws.cell(r, 3).value
            print(f"  R{r}: A={a!r} B={b!r} C={c!r} D={d_val!r}")
else:
    print("  (출력 파일 없음)")
