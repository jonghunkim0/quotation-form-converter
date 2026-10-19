"""새 템플릿 구조 점검: 섹션 행, 인건비 행, B열 병합"""
from openpyxl import load_workbook
from pathlib import Path

ws = load_workbook(Path("templates/고객사제출_표준양식.xlsx"))["2.상세 견적"]

print("=== 전체 행 (A/B/C/D/G/H/I) ===")
for r in range(1, ws.max_row + 1):
    a, b, c, d = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value
    g, h, i = ws.cell(r, 7).value, ws.cell(r, 8).value, ws.cell(r, 9).value
    if any(x not in (None, "") for x in (a, b, c, d)):
        print(f"R{r}: A={str(a)[:12]!r} B={str(b)[:10]!r} C={str(c)[:14]!r} D={str(d)[:28]!r} G={g!r} H={h!r} I={str(i)[:22]!r}")

print("\n=== B열 병합 ===")
for m in sorted(ws.merged_cells.ranges, key=lambda m: (m.min_col, m.min_row)):
    if m.min_col == 2 and m.max_col == 2:
        print(str(m))
