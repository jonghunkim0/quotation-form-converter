"""풍산 견적 파일 구조 분석 스크립트"""
import openpyxl

FILE = r"D:\견적분석\풍산 견적_AZS 각형 VER1.0(260602).xlsx"
wb = openpyxl.load_workbook(FILE, data_only=True)
ws = wb["세부"]

print("=== B열에 Unit명/소계/합계/구매품/가공품 포함 행 ===")
for r in range(1, ws.max_row + 1):
    b = ws.cell(row=r, column=2).value
    if b is not None:
        bs = str(b)
        if any(k in bs for k in ["Unit명", "소계", "합계", "구매품", "가공품",
                                  "인건비", "경비", "설계", "조립", "셋업",
                                  "양산", "포장", "관리", "이윤", "특수"]):
            a = ws.cell(row=r, column=1).value
            a_s = repr(str(a))[:20] if a else ""
            print(f"R{r}: A={a_s} | B={repr(bs)[:45]}")

print("\n=== [3]인건비 섹션 (R2160~R2175) 전체 ===")
for r in range(2160, 2176):
    vals = []
    for c in range(1, 12):
        v = ws.cell(row=r, column=c).value
        if v is not None:
            col = openpyxl.utils.get_column_letter(c)
            vals.append(f"{col}={repr(str(v))[:30]}")
    if vals:
        print(f"R{r}: " + " | ".join(vals))

print("\n=== [4]경비 섹션 (R2168~R2180) 전체 ===")
for r in range(2168, 2181):
    vals = []
    for c in range(1, 12):
        v = ws.cell(row=r, column=c).value
        if v is not None:
            col = openpyxl.utils.get_column_letter(c)
            vals.append(f"{col}={repr(str(v))[:30]}")
    if vals:
        print(f"R{r}: " + " | ".join(vals))

print("\n=== [2]전장/제어 섹션 시작 (R1809~R1830) ===")
for r in range(1809, 1831):
    vals = []
    for c in range(1, 12):
        v = ws.cell(row=r, column=c).value
        if v is not None:
            col = openpyxl.utils.get_column_letter(c)
            vals.append(f"{col}={repr(str(v))[:30]}")
    if vals:
        print(f"R{r}: " + " | ".join(vals))
