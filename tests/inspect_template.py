"""표준 양식 템플릿 구조 분석"""
import openpyxl

TPL = r"c:\AI_Agent\quotation-form-converter\templates\고객사제출_표준양식.xlsx"
wb = openpyxl.load_workbook(TPL, data_only=False)
print("=== 시트 목록 ===")
for s in wb.sheetnames:
    ws = wb[s]
    print(f"  - {s} (max_row={ws.max_row}, max_col={ws.max_column})")

print("\n=== '2.상세 견적' 시트 전체 구조 (A~K열) ===")
if "2.상세 견적" in wb.sheetnames:
    ws = wb["2.상세 견적"]
    for r in range(1, ws.max_row + 1):
        vals = []
        for c in range(1, 12):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                col = openpyxl.utils.get_column_letter(c)
                vs = str(v)
                # 수식 표시
                if vs.startswith("="):
                    vs = f"[수식]{vs[:25]}"
                vals.append(f"{col}={vs[:22]}")
        if vals:
            print(f"R{r}: " + " | ".join(vals))
