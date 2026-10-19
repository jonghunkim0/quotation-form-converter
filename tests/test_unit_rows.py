"""합산(Unit) 행 처리 검증 테스트"""
import sys
sys.path.insert(0, ".")
from converters import RuleMapper, export_standard_excel
from classification import classify_items
import openpyxl

# 원본 구조 시뮬레이션: Unit 행(No에 'Unit명') + 세부항목
test_items = [
    {'raw_data': {'No': '(1) Unit명 : BASE_FRAME', '수량': '1', '단위': '식', '단가': '31886000', '금액': '31886000'}},
    {'raw_data': {'No': '1)', '품 명': 'MAIN_FRAME', '규격': 'BT2420-L1-1001M0001-CTQ-00', 'Maker': 'PSS', '수량': '1', '단위': 'EA', '단가': '27073800', '금액': '27073800'}},
    {'raw_data': {'No': '2)', '품 명': 'FOOT_SPACER-01', '규격': 'BT2420-L1-1001M0002-00', 'Maker': 'PSS', '수량': '20', '단위': 'EA', '단가': '28100', '금액': '562000'}},
    {'raw_data': {'No': '3)', '품 명': 'FOOT_SPACER-02', '규격': 'BT2420-L1-1001M0003-00', 'Maker': 'PSS', '수량': '20', '단위': 'EA', '단가': '4600', '금액': '92000'}},
    {'raw_data': {'No': '(2) Unit명 : Z_BASE', '수량': '1', '단위': '식', '단가': '12536040', '금액': '12536040'}},
    {'raw_data': {'No': '1)', '품 명': 'Z-BASE', '규격': 'BT2420-L1-1002M0001-00', 'Maker': 'PSS', '수량': '1', '단위': 'EA', '단가': '109200', '금액': '109200'}},
    {'raw_data': {'No': '2)', '품 명': 'Z-BASE-PLATE', '규격': 'BT2420-L1-1002M0002-00', 'Maker': 'PSS', '수량': '4', '단위': 'EA', '단가': '109800', '금액': '439200'}},
]
mapper = RuleMapper()
mapped = mapper.map_items(test_items)
mapped = classify_items(mapped)
for m in mapped:
    print(f"합산={m.get('_합산')} unit={m.get('_unit_name')!r} 품명={m.get('품명')!r} -> {m.get('_분류')}")

out = export_standard_excel(mapped, company_name='풍산', quote_no='QT-1')
open('data/test_unit_output.xlsx', 'wb').write(out)
wb = openpyxl.load_workbook('data/test_unit_output.xlsx')
ws = wb['2.상세 견적']
print('--- merged (D~F 범위) ---')
for m in ws.merged_cells.ranges:
    if m.min_col <= 6 and m.max_col >= 4 and m.min_row >= 5:
        print(' ', m)
print('--- rows 5-14 ---')
for r in range(5, 15):
    c = ws.cell(r, 3).value
    d = ws.cell(r, 4).value
    g = ws.cell(r, 7).value
    h = ws.cell(r, 8).value
    i = ws.cell(r, 9).value
    j = ws.cell(r, 10).value
    print(f'R{r}: C={c!r} D={d!r} G={g!r} H={h!r} I={i!r} J={j!r}')

# ── 가공품 Unit 시나리오: 구분 행(빈칸) 삭제 검증 ──
print('\n=== 가공품 Unit 시나리오 (구분 행 삭제) ===')
test_items2 = [
    {'raw_data': {'No': '(1) Unit명 : FRAME_A', '수량': '1', '단위': '식', '단가': '1000000', '금액': '1000000'}},
    {'raw_data': {'No': '1)', '품 명': 'FRAME_PART-1', '규격': 'S1', 'Maker': 'PSS', '수량': '1', '단위': 'EA', '단가': '500000', '금액': '500000'}},
    {'raw_data': {'No': '2)', '품 명': 'FRAME_PART-2', '규격': 'S2', 'Maker': 'PSS', '수량': '2', '단위': 'EA', '단가': '250000', '금액': '500000'}},
    {'raw_data': {'No': '(2) Unit명 : PLATE_B', '수량': '1', '단위': '식', '단가': '2000000', '금액': '2000000'}},
    {'raw_data': {'No': '1)', '품 명': 'PLATE-1', '규격': 'P1', 'Maker': 'PSS', '수량': '1', '단위': 'EA', '단가': '1000000', '금액': '1000000'}},
    {'raw_data': {'No': '2)', '품 명': 'PLATE-2', '규격': 'P2', 'Maker': 'PSS', '수량': '2', '단위': 'EA', '단가': '500000', '금액': '1000000'}},
]
mapped2 = RuleMapper().map_items(test_items2)
mapped2 = classify_items(mapped2)
# FRAME_A → 구매품, PLATE_B → 가공품 (강제 지정)
for m in mapped2:
    if m.get('_합산') and m.get('_unit_name') == 'PLATE_B':
        m['_분류'] = '기구_가공품'
    elif not m.get('_합산') and m.get('품명', '').startswith('PLATE'):
        m['_분류'] = '기구_가공품'
    elif not m.get('_합산') and m.get('품명', '').startswith('FRAME'):
        m['_분류'] = '기구_구매품'
out2 = export_standard_excel(mapped2, company_name='풍산', quote_no='QT-2')
open('data/test_sep_output.xlsx', 'wb').write(out2)
wb2 = openpyxl.load_workbook('data/test_sep_output.xlsx')
ws2 = wb2['2.상세 견적']
print('--- rows 5-16 (구분 행 삭제 후) ---')
for r in range(5, 17):
    vals = []
    for c in range(1, 11):
        v = ws2.cell(r, c).value
        if v is not None:
            vals.append(f'{openpyxl.utils.get_column_letter(c)}={v!r}')
    print(f'R{r}:', ' | '.join(vals) or '(empty)')
# B열 '가공품' 라벨이 소계 바로 아래(구분 행 삭제 후)에 있는지 확인
print('--- B열 라벨 위치 ---')
for r in range(5, 17):
    b = ws2.cell(r, 2).value
    if b:
        print(f'R{r}: B={b!r}')

# ── Unit 분할 시나리오: 1개 Unit이 구매품+가공품 세부항목 혼합 ──
print('\n=== Unit 분할 시나리오 (구매품+가공품 혼합) ===')
test_items3 = [
    {'raw_data': {'No': '(1) Unit명 : BASE_FRAME', '수량': '1', '단위': '식', '단가': '31886000', '금액': '31886000'}},
    # 가공품 3개
    {'raw_data': {'No': '1)', '품 명': 'MAIN_FRAME', '규격': 'S1', 'Maker': 'PSS', '수량': '1', '단위': 'EA', '단가': '27073800', '금액': '27073800'}},
    {'raw_data': {'No': '2)', '품 명': 'FOOT_SPACER-01', '규격': 'S2', 'Maker': 'PSS', '수량': '20', '단위': 'EA', '단가': '28100', '금액': '562000'}},
    {'raw_data': {'No': '3)', '품 명': 'FOOT_SPACER-02', '규격': 'S3', 'Maker': 'PSS', '수량': '20', '단위': 'EA', '단가': '4600', '금액': '92000'}},
    # 구매품 2개 (HI MOUNT 등)
    {'raw_data': {'No': '4)', '품 명': 'HI MOUNT', '규격': 'JK15-02', 'Maker': '진건산업', '수량': '20', '단위': 'EA', '단가': '60500', '금액': '1210000'}},
    {'raw_data': {'No': '5)', '품 명': 'SPEED CONTROLLER', '규격': 'AS1002FS-06', 'Maker': 'SMC', '수량': '48', '단위': 'EA', '단가': '6490', '금액': '311520'}},
]
mapped3 = RuleMapper().map_items(test_items3)
mapped3 = classify_items(mapped3)
# 강제 분류: MAIN_FRAME/FOOT_SPACER → 가공품, HI MOUNT/SPEED CONTROLLER → 구매품
for m in mapped3:
    if m.get('_합산'):
        continue
    if m.get('품명', '').startswith(('MAIN_FRAME', 'FOOT_SPACER')):
        m['_분류'] = '기구_가공품'
    else:
        m['_분류'] = '기구_구매품'
out3 = export_standard_excel(mapped3, company_name='풍산', quote_no='QT-3')
open('data/test_split_output.xlsx', 'wb').write(out3)
wb3 = openpyxl.load_workbook('data/test_split_output.xlsx')
ws3 = wb3['2.상세 견적']
print('--- rows 5-16 (Unit 분할 후) ---')
for r in range(5, 17):
    vals = []
    for c in range(1, 11):
        v = ws3.cell(r, c).value
        if v is not None:
            vals.append(f'{openpyxl.utils.get_column_letter(c)}={v!r}')
    print(f'R{r}:', ' | '.join(vals) or '(empty)')
