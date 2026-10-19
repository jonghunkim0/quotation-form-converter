"""openpyxl delete_rows/insert_rows가 병합 셀을 자동 처리하는지 확인."""
from openpyxl import Workbook

def dump(ws, tag):
    print(f"--- {tag} ---")
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(3, 5)]
        print(f"R{r}: C={vals[0]!r} D={vals[1]!r}")
    print("merges:", sorted(str(m) for m in ws.merged_cells.ranges))

# ── delete_rows 테스트 ──
wb = Workbook(); ws = wb.active
ws.cell(1, 3).value = "H1"; ws.merge_cells("C1:D1")   # R1 병합 (삭제 행 위)
ws.cell(2, 3).value = "0)"; ws.cell(2, 4).value = "Eng"
ws.cell(3, 3).value = "1)"; ws.cell(3, 4).value = "특급"
ws.cell(4, 3).value = "2)"; ws.cell(4, 4).value = "고급"
ws.cell(5, 3).value = "H2"; ws.merge_cells("C5:D5")   # R5 병합 (삭제 행 아래)
dump(ws, "BEFORE delete row 3")
ws.delete_rows(3, 1)
dump(ws, "AFTER delete row 3 (no manual merge handling)")

# ── insert_rows 테스트 ──
wb2 = Workbook(); ws2 = wb2.active
ws2.cell(1, 3).value = "H1"; ws2.merge_cells("C1:D1")
ws2.cell(2, 3).value = "0)"; ws2.cell(2, 4).value = "Eng"
ws2.cell(3, 3).value = "1)"; ws2.cell(3, 4).value = "특급"
ws2.cell(4, 3).value = "H2"; ws2.merge_cells("C4:D4")
dump(ws2, "BEFORE insert 2 rows at row 3")
ws2.insert_rows(3, 2)
dump(ws2, "AFTER insert 2 rows at row 3 (no manual merge handling)")
