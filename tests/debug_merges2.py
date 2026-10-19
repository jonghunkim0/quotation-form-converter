"""템플릿 vs 출력 인건비 영역 병합 비교."""
from openpyxl import load_workbook

tpl = load_workbook(r"templates\고객사제출_표준양식.xlsx")["2.상세 견적"]
out = load_workbook(r"data\output\_labor_check.xlsx")["2.상세 견적"]

print("=== 템플릿 인건비 병합 (R34~R77) ===")
for m in sorted(tpl.merged_cells.ranges, key=lambda x: (x.min_row, x.min_col)):
    if 34 <= m.min_row <= 77:
        print(f"  {m}")

print("\n=== 출력 인건비 병합 (R2151~R2194) ===")
for m in sorted(out.merged_cells.ranges, key=lambda x: (x.min_row, x.min_col)):
    if 2151 <= m.min_row <= 2194:
        print(f"  {m}")

# 출력에서 C열:I열까지 넓게 병합된 것 (이상한 것)
print("\n=== 출력: C열 시작, 3열 이상 폭의 병합 (인건비 영역) ===")
for m in sorted(out.merged_cells.ranges, key=lambda x: (x.min_row, x.min_col)):
    if 2151 <= m.min_row <= 2194 and m.min_col == 3 and (m.max_col - m.min_col) >= 2:
        print(f"  {m}  (폭 {m.max_col-m.min_col+1}열)")
