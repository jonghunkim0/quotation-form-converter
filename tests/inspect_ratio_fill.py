"""비율조정 로직.xlsx 셀 채움색 분석 (노란색=입력, 회색=자동)."""
from openpyxl import load_workbook

PATH = r"D:\0_AX 가속화 팀\팀과제\재료비 자동산출\비율조정 로직.xlsx"
wb = load_workbook(PATH)
ws = wb["Sheet1"]

def fill_info(cell):
    f = cell.fill
    if f is None or f.patternType is None:
        return None
    if f.patternType == "solid":
        c = f.start_color
        if c is not None and c.type == "rgb" and c.rgb and c.rgb != "00000000":
            return str(c.rgb)
        if c is not None and c.type == "theme":
            return f"theme{c.theme}"
    return None

print("=== 채움색 있는 셀 (R1~R12 요약 테이블) ===")
for r in range(1, 13):
    for c in range(1, 17):
        cell = ws.cell(r, c)
        fi = fill_info(cell)
        if fi:
            print(f"  {cell.coordinate} fill={fi} val={cell.value!r}")

print("\n=== 채움색 있는 셀 (R14~R57 인건비 테이블, 샘플) ===")
for r in list(range(14, 18)) + list(range(29, 31)):
    for c in range(1, 17):
        cell = ws.cell(r, c)
        fi = fill_info(cell)
        if fi:
            print(f"  {cell.coordinate} fill={fi} val={cell.value!r}")

# 채움색 종류 집계
from collections import Counter
cnt = Counter()
for row in ws.iter_rows():
    for cell in row:
        fi = fill_info(cell)
        if fi:
            cnt[fi] += 1
print("\n=== 채움색 집계 ===")
for k, v in cnt.most_common():
    print(f"  {k}: {v}개")
