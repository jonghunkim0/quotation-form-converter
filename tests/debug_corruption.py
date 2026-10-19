"""파일 손상 원인 진단: externalReferences 형식, XML 무결성, 병합 유효성."""
import re
import zipfile
import sys
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parent.parent
TPL = ROOT / "templates" / "고객사제출_표준양식.xlsx"
OUT = ROOT / "data" / "output" / "_labor_check.xlsx"

print("=== 1) 템플릿 workbook.xml의 externalReferences 원본 형식 ===")
with zipfile.ZipFile(TPL) as z:
    wbxml = z.read("xl/workbook.xml").decode("utf-8")
m = re.search(r"<externalReferences>.{0,300}", wbxml, re.DOTALL)
print(repr(m.group(0)[:300]) if m else "externalReferences 없음")
# self-closing 여부
print("self-closing:", bool(re.search(r"<externalReferences\s*/>", wbxml)))
# definedName 중 [N] 포함 개수
dn = re.findall(r"<definedName[^>]*>[^<]*\[\d+\][^<]*</definedName>", wbxml)
print(f"definedName [N] 포함: {len(dn)}개")
# [N] 참조가 definedName 외에 다른 곳에 있는지
others = re.findall(r"[^>]*\[\d+\][^<]*", wbxml)
print(f"workbook.xml 전체 [N] 참조: {len(others)}개")

print("\n=== 2) 출력 파일 XML 파싱 검증 ===")
with zipfile.ZipFile(OUT) as z:
    bad = []
    for n in z.namelist():
        if n.endswith(".xml") or n.endswith(".rels"):
            try:
                ET.fromstring(z.read(n))
            except Exception as e:
                bad.append((n, str(e)))
    print("XML 파싱 실패:", bad if bad else "없음")

print("\n=== 3) 출력 파일 병합 유효성 (중복/역전/단일셀) ===")
from openpyxl import load_workbook
wb = load_workbook(OUT)
for ws in wb.worksheets:
    seen = {}
    for m in ws.merged_cells.ranges:
        key = str(m)
        if m.min_row > m.max_row or m.min_col > m.max_col:
            print(f"  [역전] {ws.title} {m}")
        if m.min_row == m.max_row and m.min_col == m.max_col:
            print(f"  [단일셀] {ws.title} {m}")
        if key in seen:
            print(f"  [중복] {ws.title} {m}")
        seen[key] = 1
print("병합 검사 완료")

print("\n=== 4) 출력 workbook.xml rels 정합성 ===")
with zipfile.ZipFile(OUT) as z:
    wbxml = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rel_ids = set(re.findall(r'Id="([^"]+)"', rels))
    used_ids = set(re.findall(r'r:id="([^"]+)"', wbxml))
    dangling = used_ids - rel_ids
    print("dangling r:id:", dangling if dangling else "없음")
    print("externalReferences 잔존:", "externalReferences" in wbxml)
    print("definedName [N] 잔존:", len(re.findall(r"\[\d+\]", wbxml)))

print("\n=== 5) C1820 병합 확인 ===")
ws = wb["2.상세 견적"]
for m in sorted(ws.merged_cells.ranges, key=lambda x: (x.min_row, x.min_col)):
    if 1815 <= m.min_row <= 1825:
        print(f"  {m}")
# 템플릿 대응 행 (R20~R22) 병합
tpl = load_workbook(TPL)["2.상세 견적"]
print("템플릿 R19~R23 병합:")
for m in sorted(tpl.merged_cells.ranges, key=lambda x: (x.min_row, x.min_col)):
    if 19 <= m.min_row <= 23:
        print(f"  {m}")
