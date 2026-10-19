"""실제 풍산 파일로 인건비 변환 전체 파이프라인 검증.
app.py를 import하면 Streamlit이 실행되므로 AST로 필요한 함수만 추출."""
import sys, ast, json
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

# ── app.py에서 필요한 함수/상수 추출 ──
src = (ROOT / "app.py").read_text(encoding="utf-8")
tree = ast.parse(src)
WANT = {"_ANALYZER_TO_SECTION", "_apply_analyzer_classification",
        "_adjust_family_by_gubun", "_reinherit_unit_classification",
        "_collect_labor_data"}
ns = {"__name__": "app_extract"}
for node in tree.body:
    if isinstance(node, ast.Assign) and any(
            isinstance(t, ast.Name) and t.id in WANT for t in node.targets):
        exec(compile(ast.Module([node], []), "app.py", "exec"), ns)
    elif isinstance(node, ast.FunctionDef) and node.name in WANT:
        exec(compile(ast.Module([node], []), "app.py", "exec"), ns)

from parsers import parse_excel
from converters import RuleMapper, export_standard_excel
from classification import classify_items
from openpyxl import load_workbook

RAW = r"D:\견적분석\풍산 견적_AZS 각형 VER1.0(260602).xlsx"

# ── Step 1-2: 파싱 + 매핑 ──
with open(RAW, "rb") as _f:
    parsed_items = parse_excel(_f)
mapper = RuleMapper({})
mapped_items = mapper.map_items(parsed_items)

# ── Step 3: analyzer 파이프라인 (app.py와 동일) ──
from quotation_analyzer import ExcelParser, AutoClassifier
a_parser = ExcelParser()
sheets = a_parser.list_sheets(RAW)
analyzer_qf = a_parser.parse_sheet(RAW, sheets[0], Path(RAW).name, "견적가")
a_classifier = AutoClassifier()
a_classifier.classify_all(analyzer_qf.items)
for it in analyzer_qf.items:
    it.final_category = it.auto_category

# ── Step 4: 분류 ──
mapped_items = classify_items(mapped_items)
ns["_apply_analyzer_classification"](mapped_items, analyzer_qf)
ns["_adjust_family_by_gubun"](mapped_items)
ns["_reinherit_unit_classification"](mapped_items)
labor_data = ns["_collect_labor_data"](analyzer_qf)
print("=== labor_data ===")
print(json.dumps(labor_data, ensure_ascii=False, indent=1))

# ── Step 5: Excel 생성 ──
out = export_standard_excel(
    mapped_items, company_name="풍산", quote_no="TEST",
    template_file=str(ROOT / "templates" / "고객사제출_표준양식.xlsx"),
    labor_data=labor_data,
)
out_path = ROOT / "data" / "output" / "_labor_check.xlsx"
out_path.write_bytes(out)
print(f"\n출력: {out_path}")

# ── 검증: 인건비 영역 행 출력 ──
wb = load_workbook(out_path)
ws = wb["2.상세 견적"]
print("\n=== 인건비 영역 (D/G/I/K) ===")
for r in range(1, ws.max_row + 1):
    d = ws.cell(r, 4).value
    c = ws.cell(r, 3).value
    if d is None and c is None:
        continue
    ds = str(d) if d is not None else ""
    if any(k in ds for k in ("Engineering", "설계자", "조립(", "엔지니어", "합계")):
        print(f"R{r}: D={ds!r} G={ws.cell(r,7).value!r} I={ws.cell(r,9).value!r} K={ws.cell(r,11).value!r}")
