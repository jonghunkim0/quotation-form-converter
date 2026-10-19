"""포장/국내 운송 Unit 원본 + 출력 위치 확인"""
import sys
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from openpyxl import load_workbook

# 원본 파일
RAW = r"D:\견적분석\풍산 견적_AZS 각형 VER1.0(260602).xlsx"
wb = load_workbook(RAW, data_only=True)
ws = wb["세부"]
print("=== 원본 파일 (포장/운송 주변) ===")
for r in range(2160, 2185):
    vals = [ws.cell(r, c).value for c in range(1, 12)]
    if any(v not in (None, "") for v in vals):
        print(f"R{r}: " + " | ".join(str(v)[:18] for v in vals if v not in (None, "")))

# 출력 파일
print("\n=== 출력 파일 (제어부_구매품 영역) ===")
ws2 = load_workbook(str(ROOT / "data/output/_labor_check.xlsx"))["2.상세 견적"]
for r in range(1800, 1830):
    c = ws2.cell(r, 3).value
    d = ws2.cell(r, 4).value
    g = ws2.cell(r, 7).value
    i = ws2.cell(r, 9).value
    if any(x not in (None, "") for x in (c, d)):
        print(f"R{r}: C={c!r} D={str(d)[:35]!r} G={g!r} I={i!r}")
