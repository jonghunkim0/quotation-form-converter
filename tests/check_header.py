"""템플릿 헤더 영역 상세 확인"""
import openpyxl

wb = openpyxl.load_workbook(r'templates/고객사제출_표준양식.xlsx')
ws = wb['1.표준 견적서(갑지)']

print('=== 이미지/도형 상세 ===')
print(f'이미지: {len(ws._images)}개')
for i, img in enumerate(ws._images):
    print(f'  [{i}] type={type(img).__name__}')
    if hasattr(img, 'path'):
        print(f'      path={img.path}')
    if hasattr(img, 'anchor'):
        anchor = img.anchor
        print(f'      anchor type={type(anchor).__name__}')
        if hasattr(anchor, '_from'):
            print(f'      from: col={anchor._from.col}, row={anchor._from.row}')
        if hasattr(anchor, 'to'):
            print(f'      to: col={anchor.to.col}, row={anchor.to.row}')

print(f'\n차트: {len(ws._charts)}개')

# Row 1~7의 모든 셀 확인
print('\n=== Row 1~7 전체 ===')
for row_idx in range(1, 8):
    for col_idx in range(1, 10):
        cell = ws.cell(row=row_idx, column=col_idx)
        v = cell.value
        if v is not None:
            print(f'  {cell.column_letter}{row_idx} = {repr(str(v)[:50])}')

# 출력 파일과 비교
print('\n\n=== 출력 파일 (test_output.xlsx) ===')
wb2 = openpyxl.load_workbook('data/test_output.xlsx')
ws2 = wb2['1.표준 견적서(갑지)']
print(f'이미지: {len(ws2._images)}개')
for i, img in enumerate(ws2._images):
    print(f'  [{i}] type={type(img).__name__}')
    if hasattr(img, 'path'):
        print(f'      path={img.path}')

print('\nRow 1~7:')
for row_idx in range(1, 8):
    for col_idx in range(1, 10):
        cell = ws2.cell(row=row_idx, column=col_idx)
        v = cell.value
        if v is not None:
            print(f'  {cell.column_letter}{row_idx} = {repr(str(v)[:50])}')
