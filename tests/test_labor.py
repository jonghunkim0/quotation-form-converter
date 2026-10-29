"""인건비 분류 규칙 테스트: 특급/고급 MD 분배 + 일액 정수 + Engineering 배수"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook
from converters.standard_exporter import _fill_detail_sheet, _fill_labor_rows

TPL = Path(__file__).resolve().parent.parent / "templates" / "고객사제출_표준양식.xlsx"

# 실제 풍산 파일 인건비 데이터 (analyzer_export.csv 기준)
LABOR = {
    "인건비-기구 설계":       {"md": 30,  "rate": 400000, "amount": 12000000, "remark": ""},
    "인건비-전장 설계":       {"md": 20,  "rate": 450000, "amount": 9000000,  "remark": ""},
    "인건비-제어 설계":       {"md": 45,  "rate": 450000, "amount": 20250000, "remark": "1名x45日(기 자료 고융 기준)"},
    "인건비-기구 조립":       {"md": 140, "rate": 400000, "amount": 56000000, "remark": "4名x20日/2名x30日"},
    "인건비-전장/제어 셋업":  {"md": 1,   "rate": 40000000, "amount": 40000000, "remark": ""},
    "인건비-기구 셋업":       {"md": 0,   "rate": 400000, "amount": 0, "remark": "해외 셋업 제외"},
    "인건비-전장/제어 양산대응": {"md": 0, "rate": 450000, "amount": 0, "remark": "해외 셋업 제외"},
}

def get(ws, r, c):
    return ws.cell(r, c).value

def check(label, actual, expected):
    ok = actual == expected
    print(f"  {'PASS' if ok else 'FAIL'} {label}: {actual!r} (expected {expected!r})")
    return ok

def main():
    wb = load_workbook(TPL)
    ws = wb["2.상세 견적"]
    # 재료비 항목 없음 → shift=0
    notes = _fill_labor_rows(ws, LABOR, 0)
    for r, text in notes:
        ws.cell(r, 13).value = text  # M열(검증결과)

    all_ok = True
    print("=== 인건비 특급/고급 분배 (인원×일수) ===")
    # 기구 설계: MD=30, 비고 없음 → 1명×30일 → 특급 1명×30일=30 / 고급 0명×30일=0
    all_ok &= check("기구설계자(특급) G", get(ws, 36, 7), 30)
    all_ok &= check("기구설계자(특급) I", get(ws, 36, 9), 450000)
    all_ok &= check("기구설계자(고급) G", get(ws, 37, 7), 0)
    all_ok &= check("기구설계자(고급) I", get(ws, 37, 9), 420000)

    # 전장 설계: MD=20, 비고 없음 → 7/14/21/30/60 후보 × 인원 탐색 → 3명×7일=21(가장 근접) → 특급=ceil(3/2)=2명×7일=14 / 고급=1명×7일=7
    all_ok &= check("전장설계자(특급) G", get(ws, 40, 7), 14)
    all_ok &= check("전장설계자(특급) I", get(ws, 40, 9), 450000)
    all_ok &= check("전장설계자(고급) G", get(ws, 41, 7), 7)
    all_ok &= check("전장설계자(고급) I", get(ws, 41, 9), 420000)

    # 제어 설계: MD=45, 비고 '1名x45日' → 단일 표현: 1명×45일 → 특급 1명×45일=45 / 고급 0
    all_ok &= check("SW설계자(특급) G", get(ws, 44, 7), 45)
    all_ok &= check("SW설계자(특급) I", get(ws, 44, 9), 450000)
    all_ok &= check("SW설계자(특급) K", get(ws, 44, 11), "1名x45日(기 자료 고융 기준)")
    # SW설계자(무관)은 단일행(single) → 고급 행 없음 (템플릿에 1행만 존재, R45는 조립 섹션 헤더)

    # 기구 조립: MD=140, 비고 '4名x20日/2名x30日' → 앞=특급 4명×20일=80 / 뒤=고급 2명×30일=60 (2026-10-29 변경)
    all_ok &= check("기구조립(특급) G", get(ws, 46, 7), 80)
    all_ok &= check("기구조립(특급) I", get(ws, 46, 9), 450000)
    all_ok &= check("기구조립(고급) G", get(ws, 47, 7), 60)
    all_ok &= check("기구조립(고급) I", get(ws, 47, 9), 420000)
    all_ok &= check("기구조립(특급) K", get(ws, 46, 11), "4名x20日/2名x30日")

    # 전장/제어 셋업: MD=1, 비고 없음 → 1명×7일 → 특급=ceil(1/2)=1명×7일=7 / 고급=0명×7일=0
    all_ok &= check("셋업제어(특급) G", get(ws, 59, 7), 7)
    all_ok &= check("셋업제어(특급) I", get(ws, 59, 9), 450000)
    all_ok &= check("셋업제어(고급) G", get(ws, 60, 7), 0)
    all_ok &= check("셋업제어(고급) I", get(ws, 60, 9), 420000)

    # 기구 셋업: MD=0, 금액=0 → 템플릿 유지 (G=0)
    all_ok &= check("셋업기구(특급) G (0유지)", get(ws, 55, 7), 0)

    # Engineering: 섹션별 최대 대응일수 합산 × 2명 (UI 방식, 2026-10-29 변경)
    # 설계 max(30,7,45)=45, 조립 max(20,30)=30, 셋업 max(7)=7, 양산대응 0 → (45+30+7)×2=164
    all_ok &= check("Engineering G", get(ws, 35, 7), 164)

    # Engineering: MD=15 → 2명×7일(14)로 재현 → 설계 max=7 → 7×2=14 (새 워크북으로 재실행)
    wb2 = load_workbook(TPL)
    ws2 = wb2["2.상세 견적"]
    _fill_labor_rows(ws2, {"인건비-기구 설계": {"md": 15, "rate": 400000, "amount": 6000000, "remark": ""}}, 0)
    all_ok &= check("Engineering G (15일→14)", get(ws2, 35, 7), 14)
    del wb2, ws2

    # J열 수식 유지
    all_ok &= check("기구설계자(특급) J 수식", get(ws, 36, 10), "=G36*I36")

    print("=== M열(검증결과) 인원×일수 표기 ===")
    all_ok &= check("M36 (기구설계)", get(ws, 36, 13), "MD=30 (1명×30일+0명×30일)")
    all_ok &= check("M44 (제어설계)", get(ws, 44, 13), "MD=45 (1명×45일+0명×45일)")
    all_ok &= check("M46 (기구조립)", get(ws, 46, 13), "MD=140 (4명×20일+2명×30일)")

    print()
    print("ALL PASS" if all_ok else "SOME FAILED")
    return 0 if all_ok else 1

if __name__ == "__main__":
    sys.exit(main())
